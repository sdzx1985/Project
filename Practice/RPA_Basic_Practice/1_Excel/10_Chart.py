from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference, LineChart

wb = load_workbook("sample.xlsx")
ws = wb.active 

# # B2:C11
# bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)
# bar_chart = BarChart() # Chart type setting
# bar_chart.add_data(bar_value) # add chart data
# ws.add_chart(bar_chart, "E1") # location the chart

# B1:C11
line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart()
line_chart.add_data(line_value, titles_from_data=True)

line_chart.title = "Grade"
line_chart.style = 15 # chart style
line_chart.y_axis.title = "Grade"
line_chart.x_axis.title = "Number"
ws.add_chart(line_chart, "E1")

wb.save("sample_chart.xlsx")